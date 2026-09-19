from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from django.conf import settings
from django.db import transaction
from django.utils import timezone

from catalog.models import ProductVariant, StockStatus
from commerce.crypto import encrypt_code, fingerprint_code, mask_code
from commerce.models import (
    CodeImportBatch,
    CodeImportBatchStatus,
    CodeImportMode,
    DigitalCode,
    DigitalCodeSource,
    DigitalCodeStatus,
)

MAX_ROWS = 10_000
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_CODE_LEN = 512
MAX_PIN_LEN = 64
PREVIEW_LIMIT = 50
ALLOWED_EXTENSIONS = {".csv", ".xlsx"}


class CodeImportError(Exception):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.message = message
        self.status = status


@dataclass
class ParsedRow:
    row_number: int
    code: str
    pin: str | None
    sku: str | None = None
    variant_id: int | None = None
    fingerprint: str = ""
    variant: ProductVariant | None = None
    error: str | None = None


@dataclass
class ValidationResult:
    rows: list[ParsedRow] = field(default_factory=list)
    total_rows: int = 0
    valid_rows: int = 0
    duplicate_rows: int = 0
    failed_rows: int = 0
    error_report: list[dict[str, Any]] = field(default_factory=list)
    preview_rows: list[dict[str, Any]] = field(default_factory=list)


def _imports_dir() -> Path:
    path = Path(settings.MEDIA_ROOT) / "code_imports"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _ext(filename: str) -> str:
    return Path(filename or "").suffix.lower()


def _normalize_header(value: str) -> str:
    return (value or "").strip().lower().replace(" ", "_")


def _encode_payload(code: str, pin: str | None) -> str:
    if pin:
        return json.dumps({"code": code, "pin": pin}, separators=(",", ":"), ensure_ascii=False)
    return code


def sync_variant_stock(variant_ids: set[int] | list[int]) -> None:
    from django.db.models import Count

    ids = {int(v) for v in variant_ids if v}
    if not ids:
        return
    available_map = {
        row["variant_id"]: row["c"]
        for row in DigitalCode.objects.filter(
            variant_id__in=ids, status=DigitalCodeStatus.AVAILABLE
        )
        .values("variant_id")
        .annotate(c=Count("id"))
    }
    for vid in ids:
        available = available_map.get(vid, 0)
        status = StockStatus.IN_STOCK if available > 0 else StockStatus.OUT_OF_STOCK
        ProductVariant.objects.filter(pk=vid).update(stock_status=status)


def parse_upload_bytes(filename: str, raw: bytes) -> list[dict[str, str]]:
    if len(raw) > MAX_FILE_BYTES:
        raise CodeImportError(f"File exceeds {MAX_FILE_BYTES // (1024 * 1024)}MB limit")
    ext = _ext(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise CodeImportError("Unsupported file type. Use .csv or .xlsx")
    if not raw:
        raise CodeImportError("File is empty")

    if ext == ".csv":
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise CodeImportError("Invalid file encoding. Use UTF-8 CSV.") from exc
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise CodeImportError("CSV has no header row")
        headers = [_normalize_header(h) for h in reader.fieldnames]
        if "code" not in headers:
            raise CodeImportError('Missing required column "code"')
        unsupported = [h for h in headers if h and h not in {"code", "pin", "sku", "variant_id"}]
        if unsupported:
            raise CodeImportError(f"Unsupported columns: {', '.join(unsupported)}")
        rows: list[dict[str, str]] = []
        for i, row in enumerate(reader, start=2):
            normalized = {_normalize_header(k): (v or "").strip() for k, v in row.items() if k is not None}
            rows.append(normalized)
            if len(rows) > MAX_ROWS:
                raise CodeImportError(f"File exceeds maximum of {MAX_ROWS:,} rows")
        return rows

    # XLSX
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise CodeImportError("XLSX support is unavailable on this server") from exc
    try:
        wb = load_workbook(filename=io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise CodeImportError("Could not read XLSX file") from exc
    ws = wb.active
    iterator = ws.iter_rows(values_only=True)
    try:
        header_cells = next(iterator)
    except StopIteration as exc:
        raise CodeImportError("Spreadsheet is empty") from exc
    headers = [_normalize_header(str(c) if c is not None else "") for c in header_cells]
    if "code" not in headers:
        raise CodeImportError('Missing required column "code"')
    unsupported = [h for h in headers if h and h not in {"code", "pin", "sku", "variant_id"}]
    if unsupported:
        raise CodeImportError(f"Unsupported columns: {', '.join(unsupported)}")
    rows = []
    for row_cells in iterator:
        if row_cells is None or all(c is None or str(c).strip() == "" for c in row_cells):
            continue
        normalized = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            cell = row_cells[idx] if idx < len(row_cells) else None
            normalized[header] = "" if cell is None else str(cell).strip()
        rows.append(normalized)
        if len(rows) > MAX_ROWS:
            raise CodeImportError(f"File exceeds maximum of {MAX_ROWS:,} rows")
    return rows


def validate_rows(
    *,
    mode: str,
    raw_rows: list[dict[str, str]],
    variant: ProductVariant | None,
) -> ValidationResult:
    result = ValidationResult(total_rows=len(raw_rows))
    if result.total_rows == 0:
        raise CodeImportError("No data rows found")

    # Resolve variants for advanced mode in bulk
    skus = {r.get("sku", "").strip() for r in raw_rows if r.get("sku")}
    variant_ids = set()
    for r in raw_rows:
        vid = (r.get("variant_id") or "").strip()
        if vid.isdigit():
            variant_ids.add(int(vid))

    sku_map = {
        v.sku: v
        for v in ProductVariant.objects.filter(sku__in=skus).select_related("product")
    } if skus else {}
    id_map = {
        v.id: v
        for v in ProductVariant.objects.filter(id__in=variant_ids).select_related("product")
    } if variant_ids else {}

    file_fingerprints: dict[str, int] = {}
    parsed: list[ParsedRow] = []

    for idx, raw in enumerate(raw_rows, start=2):
        code = (raw.get("code") or "").strip()
        pin_raw = (raw.get("pin") or "").strip()
        pin = pin_raw or None
        sku = (raw.get("sku") or "").strip() or None
        vid_raw = (raw.get("variant_id") or "").strip()
        row = ParsedRow(row_number=idx, code=code, pin=pin, sku=sku)

        if not code:
            row.error = "Code is empty"
        elif len(code) > MAX_CODE_LEN:
            row.error = "Code exceeds maximum length"
        elif pin and len(pin) > MAX_PIN_LEN:
            row.error = "PIN exceeds maximum length"
        else:
            if mode == CodeImportMode.SIMPLE:
                if not variant:
                    row.error = "Variant is required for simple mode"
                else:
                    row.variant = variant
                    row.variant_id = variant.id
            else:
                if sku and sku in sku_map:
                    row.variant = sku_map[sku]
                    row.variant_id = row.variant.id
                elif vid_raw.isdigit() and int(vid_raw) in id_map:
                    row.variant = id_map[int(vid_raw)]
                    row.variant_id = row.variant.id
                    row.sku = row.variant.sku
                elif sku:
                    row.error = f'SKU "{sku}" not found'
                elif vid_raw:
                    row.error = f'Variant id "{vid_raw}" not found'
                else:
                    row.error = "Row must include sku or variant_id"

        if not row.error:
            try:
                row.fingerprint = fingerprint_code(code, pin)
            except Exception:
                row.error = "Could not fingerprint code"
            else:
                if row.fingerprint in file_fingerprints:
                    row.error = f"Duplicate code in file (same as row {file_fingerprints[row.fingerprint]})"
                    result.duplicate_rows += 1
                else:
                    file_fingerprints[row.fingerprint] = idx

        parsed.append(row)

    # DB duplicate check in bulk
    fps = [r.fingerprint for r in parsed if r.fingerprint and not r.error]
    existing: set[str] = set()
    if fps:
        # Chunk to avoid oversized IN clauses
        chunk_size = 2000
        for i in range(0, len(fps), chunk_size):
            chunk = fps[i : i + chunk_size]
            existing.update(
                DigitalCode.objects.filter(code_fingerprint__in=chunk).values_list(
                    "code_fingerprint", flat=True
                )
            )

    for row in parsed:
        if row.error:
            result.failed_rows += 1
            result.error_report.append({"row": row.row_number, "error": row.error, "sku": row.sku or ""})
        elif row.fingerprint in existing:
            row.error = "Duplicate code already exists"
            result.duplicate_rows += 1
            result.failed_rows += 1
            result.error_report.append({"row": row.row_number, "error": row.error, "sku": row.sku or ""})
        else:
            result.valid_rows += 1

        if len(result.preview_rows) < PREVIEW_LIMIT:
            result.preview_rows.append(
                {
                    "row": row.row_number,
                    "sku": row.sku or (row.variant.sku if row.variant else ""),
                    "variant_id": row.variant_id,
                    "code_masked": mask_code(row.code) if row.code else "",
                    "pin_masked": mask_code(row.pin) if row.pin else "",
                    "status": "ok" if not row.error else "error",
                    "error": row.error or "",
                }
            )

    result.rows = parsed
    return result


def preview_import(
    *,
    uploaded_by,
    filename: str,
    raw: bytes,
    mode: str,
    variant_id: int | None,
) -> CodeImportBatch:
    if mode not in {CodeImportMode.SIMPLE, CodeImportMode.ADVANCED}:
        raise CodeImportError("Invalid import mode")
    variant = None
    if mode == CodeImportMode.SIMPLE:
        if not variant_id:
            raise CodeImportError("variant_id is required for simple mode")
        variant = ProductVariant.objects.filter(pk=variant_id).select_related("product").first()
        if not variant:
            raise CodeImportError("Variant not found", status=404)

    raw_rows = parse_upload_bytes(filename, raw)
    validation = validate_rows(mode=mode, raw_rows=raw_rows, variant=variant)

    token = uuid.uuid4().hex
    ext = _ext(filename) or ".csv"
    rel_path = f"code_imports/{token}{ext}"
    abs_path = _imports_dir() / f"{token}{ext}"
    abs_path.write_bytes(raw)

    batch = CodeImportBatch.objects.create(
        filename=os.path.basename(filename)[:255],
        mode=mode,
        uploaded_by=uploaded_by,
        variant=variant,
        status=CodeImportBatchStatus.VALIDATED if validation.valid_rows else CodeImportBatchStatus.FAILED,
        total_rows=validation.total_rows,
        valid_rows=validation.valid_rows,
        imported_rows=0,
        duplicate_rows=validation.duplicate_rows,
        failed_rows=validation.failed_rows,
        error_report=validation.error_report,
        preview_rows=validation.preview_rows,
        pending_file_path=rel_path,
        file_sha256=hashlib.sha256(raw).hexdigest(),
    )
    return batch


def confirm_import(*, batch_id: int, actor) -> CodeImportBatch:
    batch = CodeImportBatch.objects.select_related("variant", "uploaded_by").filter(pk=batch_id).first()
    if not batch:
        raise CodeImportError("Import batch not found", status=404)
    if batch.imported_rows > 0 and batch.status in {
        CodeImportBatchStatus.COMPLETED,
        CodeImportBatchStatus.PARTIAL,
    }:
        raise CodeImportError("Import batch already completed", status=409)
    if batch.status not in {CodeImportBatchStatus.VALIDATED}:
        if batch.status == CodeImportBatchStatus.FAILED and batch.valid_rows == 0:
            raise CodeImportError("No valid rows to import")
        raise CodeImportError(f"Batch cannot be confirmed (status={batch.status})")
    if not batch.pending_file_path:
        raise CodeImportError("Pending upload file is missing; re-upload for preview")

    abs_path = Path(settings.MEDIA_ROOT) / batch.pending_file_path
    if not abs_path.is_file():
        raise CodeImportError("Pending upload file expired; re-upload for preview")

    raw = abs_path.read_bytes()
    if hashlib.sha256(raw).hexdigest() != batch.file_sha256:
        raise CodeImportError("Uploaded file changed; re-run preview")

    raw_rows = parse_upload_bytes(batch.filename, raw)
    validation = validate_rows(mode=batch.mode, raw_rows=raw_rows, variant=batch.variant)

    batch.status = CodeImportBatchStatus.PROCESSING
    batch.save(update_fields=["status"])

    to_create: list[DigitalCode] = []
    variant_ids: set[int] = set()
    for row in validation.rows:
        if row.error or not row.variant:
            continue
        payload = _encode_payload(row.code, row.pin)
        ciphertext, nonce, _masked = encrypt_code(payload)
        to_create.append(
            DigitalCode(
                variant=row.variant,
                status=DigitalCodeStatus.AVAILABLE,
                source=DigitalCodeSource.CSV_IMPORT,
                code_fingerprint=row.fingerprint,
                import_batch=batch,
                imported_by=actor,
                ciphertext=ciphertext,
                nonce=nonce,
                masked=mask_code(row.code),
            )
        )
        variant_ids.add(row.variant.id)

    with transaction.atomic():
        if to_create:
            chunk_size = 500
            for i in range(0, len(to_create), chunk_size):
                DigitalCode.objects.bulk_create(to_create[i : i + chunk_size], ignore_conflicts=True)

        imported = DigitalCode.objects.filter(import_batch=batch).count()

        if imported > 0 and validation.failed_rows == 0 and imported >= validation.valid_rows:
            status = CodeImportBatchStatus.COMPLETED
        elif imported > 0:
            status = CodeImportBatchStatus.PARTIAL
        else:
            status = CodeImportBatchStatus.FAILED

        batch.status = status
        batch.imported_rows = imported
        batch.valid_rows = validation.valid_rows
        batch.duplicate_rows = validation.duplicate_rows
        batch.failed_rows = validation.failed_rows
        batch.error_report = validation.error_report
        batch.preview_rows = validation.preview_rows
        batch.completed_at = timezone.now()
        batch.pending_file_path = ""
        batch.save(
            update_fields=[
                "status",
                "imported_rows",
                "valid_rows",
                "duplicate_rows",
                "failed_rows",
                "error_report",
                "preview_rows",
                "completed_at",
                "pending_file_path",
            ]
        )

    try:
        abs_path.unlink(missing_ok=True)
    except OSError:
        pass

    sync_variant_stock(variant_ids)
    return batch


def create_manual_code(*, variant_id: int, code: str, pin: str | None, actor) -> DigitalCode:
    variant = ProductVariant.objects.filter(pk=variant_id).first()
    if not variant:
        raise CodeImportError("Variant not found", status=404)
    code = (code or "").strip()
    pin = (pin or "").strip() or None
    if not code:
        raise CodeImportError("Code is required")
    if len(code) > MAX_CODE_LEN:
        raise CodeImportError("Code exceeds maximum length")
    if pin and len(pin) > MAX_PIN_LEN:
        raise CodeImportError("PIN exceeds maximum length")
    fp = fingerprint_code(code, pin)
    if DigitalCode.objects.filter(code_fingerprint=fp).exists():
        raise CodeImportError("Duplicate code already exists")
    ciphertext, nonce, _ = encrypt_code(_encode_payload(code, pin))
    obj = DigitalCode.objects.create(
        variant=variant,
        status=DigitalCodeStatus.AVAILABLE,
        source=DigitalCodeSource.MANUAL,
        code_fingerprint=fp,
        imported_by=actor,
        ciphertext=ciphertext,
        nonce=nonce,
        masked=mask_code(code),
    )
    sync_variant_stock({variant.id})
    return obj


def deactivate_code(*, code_id: int, actor) -> DigitalCode:
    obj = DigitalCode.objects.filter(pk=code_id).first()
    if not obj:
        raise CodeImportError("Code not found", status=404)
    if obj.status != DigitalCodeStatus.AVAILABLE:
        raise CodeImportError("Only available unused codes can be deactivated")
    obj.status = DigitalCodeStatus.UNAVAILABLE
    obj.save(update_fields=["status"])
    if obj.variant_id:
        sync_variant_stock({obj.variant_id})
    return obj


def template_csv(mode: str) -> str:
    if mode == CodeImportMode.ADVANCED:
        return "sku,code,pin\nPSN-10-US,SAMPLE-CODE-001,1234\nPSN-20-US,SAMPLE-CODE-002,\n"
    return "code,pin\nSAMPLE-CODE-001,1234\nSAMPLE-CODE-002,\n"
